import io
import os
from datetime import datetime
from typing import Dict, Any, List
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

DEFAULT_TEMPLATE_PATH = "template.docx"

class ReportGeneratorService:
    def __init__(self, template_path: str = DEFAULT_TEMPLATE_PATH):
        self.template_path = template_path

    def generate_docx(self, study_name: str, address: str, buffer_meters: int, parcels: List[Dict[str, Any]], collected_data: Dict[str, Any]) -> io.BytesIO:
        if os.path.exists(self.template_path):
            doc = docx.Document(self.template_path)
        else:
            doc = docx.Document()

        # Styles & Colors
        title_color = RGBColor(34, 75, 45) # Deep Earth Green

        # Document Header
        title = doc.add_heading("SOIL'UP - RAPPORT DE PRÉ-DIAGNOSTIC DOCUMENTAIRE (SSP)", level=0)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in title.runs:
            run.font.color.rgb = title_color
            run.font.size = Pt(18)
            run.font.bold = True

        p_subtitle = doc.add_paragraph("Étude historique préliminaire des risques environnementaux et pollutions des sols (NF X 31-620 Phase 1)")
        p_subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        p_site = doc.add_paragraph()
        run_name = p_site.add_run(f"Site d'étude : {study_name}\n")
        run_name.bold = True
        run_name.font.size = Pt(12)
        p_site.add_run(f"Adresse : {address}\n")
        p_site.paragraph_format.space_after = Pt(20)

        # 1. Identification du site d'étude
        h1 = doc.add_heading("1. Fiche d'identification du site d'étude", level=1)
        for run in h1.runs:
            run.font.color.rgb = title_color

        meta_table = doc.add_table(rows=4, cols=2)
        meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        meta_table.style = "Table Grid"
        
        meta_data = [
            ("Nom / Référence du site", study_name),
            ("Adresse d'étude", address),
            ("Date d'analyse", datetime.now().strftime("%d/%m/%Y")),
            ("Buffer d'investigation retenu", f"{buffer_meters} m"),
        ]
        for i, (k, v) in enumerate(meta_data):
            cell_k = meta_table.cell(i, 0)
            cell_v = meta_table.cell(i, 1)
            cell_k.text = k
            cell_k.paragraphs[0].runs[0].font.bold = True
            cell_v.text = str(v)

        doc.add_paragraph().paragraph_format.space_after = Pt(15)

        # 2. Contexte Cadastral
        h2 = doc.add_heading("2. Assiette foncière et parcelles cadastrales", level=1)
        for run in h2.runs:
            run.font.color.rgb = title_color

        if parcels:
            p_table = doc.add_table(rows=1, cols=4)
            p_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            p_table.style = "Table Grid"
            headers = ["Section", "Numéro", "Contenance (m²)", "Identifiant Parcelle (IDU)"]
            for col_idx, htext in enumerate(headers):
                cell = p_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for p in parcels:
                row_cells = p_table.add_row().cells
                row_cells[0].text = str(p.get("section", "-"))
                row_cells[1].text = str(p.get("numero", "-"))
                row_cells[2].text = str(p.get("contenance", "-"))
                row_cells[3].text = str(p.get("idu", "-"))
        else:
            doc.add_paragraph("Aucune parcelle cadastrale renseignée (emprise tracée librement).")

        doc.add_paragraph().paragraph_format.space_after = Pt(15)

        # 3. Contexte Topographique, Géologique & Hydrogéologique
        h3 = doc.add_heading("3. Contexte Topographique, Géologique & Hydrogéologique", level=1)
        for run in h3.runs:
            run.font.color.rgb = title_color

        # Synthèse prioritaire : Profondeur de la Nappe au droit du site
        gw = collected_data.get("groundwater_level")
        if gw and gw.get("profondeur_estimee_m") is not None:
            p_gw_title = doc.add_paragraph()
            r_gw_title = p_gw_title.add_run("💧 ESTIMATION DU NIVEAU DE LA NAPPE AU DROIT DU SITE (NF X 31-620) :")
            r_gw_title.bold = True
            r_gw_title.font.color.rgb = title_color

            gw_table = doc.add_table(rows=4, cols=2)
            gw_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            gw_table.style = "Table Grid"

            gw_data = [
                ("Profondeur estimée sous le terrain naturel", f"~ {gw.get('profondeur_estimee_m')} m / sol"),
                ("Fourchette locale observée (hautes / basses eaux)", f"{gw.get('profondeur_min_m')} m à {gw.get('profondeur_max_m')} m / sol"),
                ("Cote piézométrique absolue estimée", f"{gw.get('cote_nappe_estimee_ngf')} m NGF" if gw.get('cote_nappe_estimee_ngf') is not None else "Non déterminée"),
                ("Fiabilité & Indice de confiance", f"{gw.get('indice_confiance')} ({gw.get('nb_points_appui')} ouvrage(s) d'appui exploité(s))"),
            ]
            for row_i, (gk, gv) in enumerate(gw_data):
                cell_k = gw_table.cell(row_i, 0)
                cell_v = gw_table.cell(row_i, 1)
                cell_k.text = gk
                cell_k.paragraphs[0].runs[0].font.bold = True
                cell_v.text = str(gv)

            p_method = doc.add_paragraph()
            p_method.paragraph_format.space_before = Pt(6)
            p_method.add_run("Méthode de calcul : ").bold = True
            p_method.add_run(str(gw.get("methode_calcul", "")))

            doc.add_paragraph()

        # Altimétrie MNT
        alti = collected_data.get("altimetrie")
        if alti and alti.get("z_mean") is not None:
            p_alti = doc.add_paragraph()
            p_alti.add_run("Topographie & Relief (IGN RGE ALTI) : ").bold = True
            p_alti.add_run(
                f"L'altitude moyenne du site est de {alti.get('z_mean')} m NGF-IGN69 "
                f"(min : {alti.get('z_min')} m, max : {alti.get('z_max')} m, dénivelé : {alti.get('denivele')} m)."
            )
        else:
            doc.add_paragraph("Données altimétriques non disponibles pour ce secteur.")

        # Hydrogéologie & Piézomètres
        piezo_list = collected_data.get("piezometrie", [])
        p_piezo_title = doc.add_paragraph()
        p_piezo_title.add_run("Hydrogéologie & Nappe Phréatique (Hub'Eau) : ").bold = True
        if piezo_list:
            closest_piezo = piezo_list[0]
            prof_str = f"{closest_piezo.get('profondeur_nappe')} m sous le sol" if closest_piezo.get('profondeur_nappe') is not None else "non mesurée"
            cote_str = f"{closest_piezo.get('cote_ngf')} m NGF" if closest_piezo.get('cote_ngf') is not None else "non renseignée"
            p_piezo_title.add_run(
                f"Au piézomètre le plus proche ({closest_piezo.get('code_bss')}, situé à {closest_piezo.get('distance_meters')} m), "
                f"le niveau de l'eau souterraine est relevé à {prof_str} (cote : {cote_str})."
            )

            piezo_table = doc.add_table(rows=1, cols=5)
            piezo_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            piezo_table.style = "Table Grid"
            pz_headers = ["Code BSS", "Commune", "Profondeur Eau (m)", "Cote Nappe (m NGF)", "Distance (m)"]
            for col_idx, htext in enumerate(pz_headers):
                cell = piezo_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for pz in piezo_list[:5]:
                row_cells = piezo_table.add_row().cells
                row_cells[0].text = str(pz.get("code_bss", "-"))
                row_cells[1].text = str(pz.get("nom_commune", "-"))
                row_cells[2].text = str(pz.get("profondeur_nappe") if pz.get("profondeur_nappe") is not None else "-")
                row_cells[3].text = str(pz.get("cote_ngf") if pz.get("cote_ngf") is not None else "-")
                row_cells[4].text = f"{pz.get('distance_meters')} m"
        else:
            p_piezo_title.add_run(f"Aucun piézomètre recensé dans le buffer de {buffer_meters} m.")

        doc.add_paragraph()

        # Sondages BSS
        bss_list = collected_data.get("bss", [])
        p_bss_title = doc.add_paragraph()
        p_bss_title.add_run("Banque du Sous-Sol & Sondages de reconnaissance (BRGM) : ").bold = True
        if bss_list:
            nb_coupes = sum(1 for b in bss_list if b.get("coupe_disponible"))
            p_bss_title.add_run(
                f"{len(bss_list)} ouvrage(s) BSS recensé(s) dans le périmètre d'étude, "
                f"dont {nb_coupes} disposant d'une coupe géologique / lithologique numérisée."
            )

            bss_table = doc.add_table(rows=1, cols=6)
            bss_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            bss_table.style = "Table Grid"
            bss_headers = ["Référence BSS", "Nature", "Prof. (m)", "Eau (m/sol)", "Coupe lithologique", "Distance (m)"]
            for col_idx, htext in enumerate(bss_headers):
                cell = bss_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for b in bss_list[:10]:
                row_cells = bss_table.add_row().cells
                row_cells[0].text = str(b.get("reference", "-"))
                row_cells[1].text = str(b.get("nature", "FORAGE"))
                row_cells[2].text = str(b.get("prof_atteinte") if b.get("prof_atteinte") is not None else "-")
                row_cells[3].text = f"{b.get('prof_eau_sol')} m" if b.get("prof_eau_sol") is not None else "Non relevée"
                row_cells[4].text = "Disponible (InfoTerre)" if b.get("coupe_disponible") else "Non numérisée"
                row_cells[5].text = f"{b.get('distance_meters')} m"
        else:
            p_bss_title.add_run(f"Aucun forage BSS recensé dans le buffer de {buffer_meters} m.")

        doc.add_paragraph().paragraph_format.space_after = Pt(15)

        # 4. Contexte Historique Industriel (BASIAS / CASIAS)
        h4 = doc.add_heading("4. Inventaire historique des activités industrielles (BASIAS / CASIAS)", level=1)
        for run in h4.runs:
            run.font.color.rgb = title_color

        casias_list = collected_data.get("casias", [])
        if casias_list:
            c_table = doc.add_table(rows=1, cols=5)
            c_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            c_table.style = "Table Grid"
            cheaders = ["Identifiant", "Nom de l'établissement", "Activité historique", "Statut", "Distance (m)"]
            for col_idx, htext in enumerate(cheaders):
                cell = c_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for item in casias_list:
                row_cells = c_table.add_row().cells
                row_cells[0].text = str(item.get("identifiant", "-"))
                row_cells[1].text = str(item.get("nom_etablissement", "-"))
                row_cells[2].text = str(item.get("activite_principale") or "Non précisée")
                row_cells[3].text = str(item.get("statut") or "-")
                dist = item.get("distance_meters", 0.0)
                row_cells[4].text = "Emprise directe" if dist == 0 else f"{dist} m"
        else:
            doc.add_paragraph(f"Aucune activité historique industrielle recensée dans le buffer de {buffer_meters} m.")

        doc.add_paragraph().paragraph_format.space_after = Pt(15)

        # 5. Installations Classées (ICPE)
        h5 = doc.add_heading("5. Installations Classées pour la Protection de l'Environnement (ICPE)", level=1)
        for run in h5.runs:
            run.font.color.rgb = title_color

        icpe_list = collected_data.get("icpe", [])
        if icpe_list:
            i_table = doc.add_table(rows=1, cols=4)
            i_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            i_table.style = "Table Grid"
            iheaders = ["Code AIOT", "Raison Sociale", "Commune", "Distance (m)"]
            for col_idx, htext in enumerate(iheaders):
                cell = i_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for item in icpe_list:
                row_cells = i_table.add_row().cells
                row_cells[0].text = str(item.get("code_aiot", "-"))
                row_cells[1].text = str(item.get("raison_sociale", "-"))
                row_cells[2].text = str(item.get("commune", "-"))
                dist = item.get("distance_meters", 0.0)
                row_cells[3].text = "Emprise directe" if dist == 0 else f"{dist} m"
        else:
            doc.add_paragraph(f"Aucune installation classée (ICPE) recensée dans le buffer de {buffer_meters} m.")

        doc.add_paragraph().paragraph_format.space_after = Pt(15)

        # 6. Sensibilité Écologique & Milieux Naturels (INPN)
        h6 = doc.add_heading("6. Sensibilité Écologique & Milieux Naturels", level=1)
        for run in h6.runs:
            run.font.color.rgb = title_color

        nature_list = collected_data.get("zonages_naturels", [])
        if nature_list:
            p_nat = doc.add_paragraph()
            p_nat.add_run(f"{len(nature_list)} zonage(s) naturel(s) et écologique(s) inventorié(s) à proximité du site :")

            n_table = doc.add_table(rows=1, cols=4)
            n_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            n_table.style = "Table Grid"
            nheaders = ["Type de Zonage", "Nom du Site", "Distance minimale", "Statut / Lien"]
            for col_idx, htext in enumerate(nheaders):
                cell = n_table.cell(0, col_idx)
                cell.text = htext
                cell.paragraphs[0].runs[0].font.bold = True

            for zn in nature_list[:10]:
                row_cells = n_table.add_row().cells
                row_cells[0].text = str(zn.get("type_zonage", "-"))
                row_cells[1].text = str(zn.get("nom_site", "-"))
                dist_m = zn.get("distance_meters", 0.0)
                row_cells[2].text = "Recoupement direct (Sur site)" if dist_m == 0 else f"{dist_m} m"
                row_cells[3].text = str(zn.get("url_fiche") or f"INPN ({zn.get('id_mnhn')})")
        else:
            doc.add_paragraph(f"Aucun zonage naturel protégé (ZNIEFF, Natura 2000) recensé dans le buffer de {buffer_meters} m.")

        # Save to BytesIO
        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)
        return buffer
