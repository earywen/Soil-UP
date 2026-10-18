import io
import pytest
import docx
from httpx import AsyncClient, ASGITransport
from soil_up.main import app

SAMPLE_EMPRISE_GEOJSON = {
    "type": "Polygon",
    "coordinates": [
        [[2.330, 48.869], [2.332, 48.869], [2.332, 48.871], [2.330, 48.871], [2.330, 48.869]]
    ]
}

@pytest.mark.asyncio
async def test_export_study_docx_returns_valid_document():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        # 1. Create a study
        create_payload = {
            "name": "Audit Foncier Ancien Dépôt",
            "address": "25 Rue du Faubourg Saint-Honoré 75008 Paris",
            "emprise": SAMPLE_EMPRISE_GEOJSON,
            "buffer_meters": 500,
            "parcels": [
                {"idu": "75108000AB0042", "section": "AB", "numero": "0042", "contenance": 850}
            ]
        }
        res_create = await client.post("/api/studies/", json=create_payload)
        assert res_create.status_code == 201
        study_id = res_create.json()["id"]

        # 2. Mock a completed collection by updating study or calling collect
        # Let's populate collected_data by updating directly or via test helper
        # We can update the study record with status completed and collected_data
        from soil_up.database import SessionLocal
        from soil_up.models.study import StudyModel
        db = SessionLocal()
        record = db.query(StudyModel).filter(StudyModel.id == study_id).first()
        record.status = "completed"
        record.collected_data = {
            "casias": [
                {
                    "identifiant": "IDF7500123",
                    "nom_etablissement": "Usine Métallurgique Saint-Honoré",
                    "activite_principale": "Traitement de surfaces",
                    "statut": "En arrêt",
                    "distance_meters": 135.5,
                    "latitude": 48.87,
                    "longitude": 2.33
                }
            ],
            "icpe": [
                {
                    "code_aiot": "0009999999",
                    "raison_sociale": "Atelier Mécanique Central",
                    "commune": "Paris",
                    "distance_meters": 240.0,
                    "latitude": 48.87,
                    "longitude": 2.33
                }
            ]
        }
        db.commit()
        db.close()

        # 3. Export Word document
        res_export = await client.get(f"/api/studies/{study_id}/export/docx")
        assert res_export.status_code == 200
        assert "application/vnd.openxmlformats-officedocument.wordprocessingml.document" in res_export.headers.get("content-type", "")
        assert "attachment" in res_export.headers.get("content-disposition", "")

        # 4. Parse binary docx content
        doc = docx.Document(io.BytesIO(res_export.content))
        full_text = " ".join([p.text for p in doc.paragraphs])
        assert "Audit Foncier Ancien Dépôt" in full_text
        assert "25 Rue du Faubourg Saint-Honoré" in full_text

        # Verify tables contain parcel and casias data
        table_texts = []
        for table in doc.tables:
            for row in table.rows:
                table_texts.append(" ".join([cell.text for cell in row.cells]))
        combined_tables = " ".join(table_texts)
        assert "75108000AB0042" in combined_tables
        assert "Usine Métallurgique Saint-Honoré" in combined_tables
        assert "135.5" in combined_tables or "135" in combined_tables

@pytest.mark.asyncio
async def test_export_study_not_collected_returns_400():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        res_create = await client.post("/api/studies/", json={
            "name": "Audit Non Collecté",
            "address": "Paris",
            "emprise": SAMPLE_EMPRISE_GEOJSON,
            "buffer_meters": 500
        })
        study_id = res_create.json()["id"]

        res_export = await client.get(f"/api/studies/{study_id}/export/docx")
        assert res_export.status_code == 400
        assert "collecte" in res_export.json()["detail"].lower()

@pytest.mark.asyncio
async def test_export_non_existent_study_returns_404():
    async with AsyncClient(transport=ASGITransport(app=app), base_url="http://test") as client:
        res_export = await client.get("/api/studies/99999/export/docx")
        assert res_export.status_code == 404
