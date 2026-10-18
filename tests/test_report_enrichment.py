import io
import pytest
import docx
from soil_up.services.report import ReportGeneratorService

SAMPLE_JALON2_COLLECTED_DATA = {
    "casias": [
        {
            "identifiant": "IDF7500123",
            "nom_etablissement": "Usine Métallurgique Saint-Honoré",
            "activite_principale": "Traitement de surfaces",
            "statut": "En arrêt",
            "distance_meters": 135.5,
        }
    ],
    "icpe": [],
    "altimetrie": {
        "z_mean": 42.5,
        "z_min": 40.2,
        "z_max": 44.8,
        "denivele": 4.6,
        "unit": "m NGF-IGN69"
    },
    "bss": [
        {
            "reference": "00994X0356/111111",
            "bss_id": "BSS000GNGE",
            "nature": "FORAGE",
            "prof_atteinte": 203.8,
            "zsol": 42.0,
            "coupe_disponible": True,
            "nb_scans_coupe": 6,
            "url_infoterre": "http://ficheinfoterre.brgm.fr/InfoterreFiche/ficheBss.action?id=00994X0356/111111",
            "distance_meters": 120.0,
            "latitude": 48.87,
            "longitude": 2.33
        }
    ],
    "piezometrie": [
        {
            "code_bss": "00994D0188/P2",
            "nom_commune": "Paris",
            "distance_meters": 350.0,
            "profondeur_nappe": 5.2,
            "cote_ngf": 36.8,
            "date_mesure": "2024-05-10"
        }
    ],
    "zonages_naturels": [
        {
            "id_mnhn": "230030439",
            "nom_site": "Bord de Seine et Îles associées",
            "type_zonage": "ZNIEFF de type I",
            "distance_meters": 450.0,
            "statut_recoupement": "À 450.0 m",
            "url_fiche": "https://inpn.mnhn.fr/zone/znieff/230030439"
        }
    ]
}

def test_report_service_enriches_docx_with_geology_and_nature():
    service = ReportGeneratorService()
    buffer = service.generate_docx(
        study_name="Site Test Jalon 2",
        address="10 Rue de la Paix 75002 Paris",
        buffer_meters=500,
        parcels=[{"section": "AX", "numero": "12", "contenance": 1200, "idu": "75102000AX0012"}],
        collected_data=SAMPLE_JALON2_COLLECTED_DATA
    )

    doc = docx.Document(buffer)
    full_text = " ".join([p.text for p in doc.paragraphs])

    # Check Topography, Geology & Hydrogeology sections
    assert "Contexte Topographique, Géologique & Hydrogéologique" in full_text
    assert "42.5 m NGF" in full_text or "42.5" in full_text
    assert "Sensibilité Écologique & Milieux Naturels" in full_text

    # Check tables content
    table_texts = []
    for table in doc.tables:
        for row in table.rows:
            table_texts.append(" ".join([cell.text for cell in row.cells]))
    combined_tables = " ".join(table_texts)

    # BSS table
    assert "00994X0356/111111" in combined_tables
    assert "203.8" in combined_tables
    # Piezo
    assert "00994D0188/P2" in combined_tables
    assert "5.2" in combined_tables
    # ZNIEFF
    assert "Bord de Seine et Îles associées" in combined_tables
    assert "ZNIEFF de type I" in combined_tables
